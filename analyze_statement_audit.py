import openpyxl

wb = openpyxl.load_workbook(r'C:\SinaankProjects\.   . Trading\SUBHMANI-BTC\Upto1910_transaction_statement.xlsx', data_only=True)

# 1. Fee Sheet
fee_sheet = wb['Future Fee history (INR-M)']
fee_rows = list(fee_sheet.iter_rows(values_only=True))
header_idx = 3
headers = [str(c).strip() if c is not None else '' for c in fee_rows[header_idx]]

fee_types = set()
total_fee_by_type_oct19 = {}
total_fee_oct19 = 0
fee_entries_oct19 = []

for r in fee_rows[header_idx + 1:]:
    row_dict = dict(zip(headers, r))
    ftype = row_dict.get('Fee Type')
    date_str = str(row_dict.get('Date (IST)'))
    amt = row_dict.get('Amount (INR)')
    fee_types.add(ftype)
    
    if 'October 19, 2026' in date_str and amt is not None:
        try:
            amt_f = float(amt)
            total_fee_by_type_oct19[ftype] = total_fee_by_type_oct19.get(ftype, 0) + amt_f
            fee_entries_oct19.append((date_str, ftype, amt_f, row_dict.get('Fee Rate')))
        except:
            pass

print("=== ALL FEE TYPES IN STATEMENT ===")
print(fee_types)
print("\n=== OCT 19, 2026 FEES BREAKDOWN ===")
for k, v in total_fee_by_type_oct19.items():
    print(f"  {k}: Rs. {v:.2f}")

# 2. Position Sheet
pos_sheet = wb['Future Position history (INR-M)']
pos_rows = list(pos_sheet.iter_rows(values_only=True))
p_headers = [str(c).strip() if c is not None else '' for c in pos_rows[6]]

oct19_positions = []
total_gross_pnl_oct19 = 0

for r in pos_rows[7:]:
    row_dict = dict(zip(p_headers, r))
    entry_time = str(row_dict.get('Entry Time (IST)'))
    if 'October 19, 2026' in entry_time:
        pnl = float(row_dict.get('Position PNL (INR)') or 0)
        total_gross_pnl_oct19 += pnl
        oct19_positions.append(row_dict)

print(f"\n=== OCT 19, 2026 POSITIONS (Total: {len(oct19_positions)}) ===")
for p in oct19_positions:
    print(f"Entry: {p['Entry Time (IST)']} | Exit: {p['Exit Time (IST)']} | Qty: {p['Quantity']} | Entry: {p['Entry Price (USDT)']} | Exit: {p['Exit Price (USDT)']} | PNL: Rs. {p['Position PNL (INR)']}")

# 3. Oct 18, 2026 (Yesterday) for comparison
oct18_positions = []
total_gross_pnl_oct18 = 0
for r in pos_rows[7:]:
    row_dict = dict(zip(p_headers, r))
    entry_time = str(row_dict.get('Entry Time (IST)'))
    if 'October 18, 2026' in entry_time:
        pnl = float(row_dict.get('Position PNL (INR)') or 0)
        total_gross_pnl_oct18 += pnl
        oct18_positions.append(row_dict)

print(f"\n=== OCT 18, 2026 POSITIONS (Total: {len(oct18_positions)}) ===")
print(f"Total Gross PnL on Oct 18: Rs. {total_gross_pnl_oct18:.2f}")

transaction_fee = total_fee_by_type_oct19.get('TRANSACTION', 0)
rebate = total_fee_by_type_oct19.get('REBATE', 0)
funding_fee = total_fee_by_type_oct19.get('FUNDING', 0)
net_fee_oct19 = transaction_fee - rebate + funding_fee

print("\n=== FINAL RECONCILIATION FOR OCT 19, 2026 ===")
print(f"Gross Position PnL:        +Rs. {total_gross_pnl_oct19:.2f}")
print(f"Transaction Fees Deducted:  -Rs. {transaction_fee:.2f}")
print(f"Rebates Credited:           +Rs. {rebate:.2f}")
print(f"Funding Fees:                Rs. {funding_fee:.2f}")
print(f"--------------------------------------------------")
print(f"Net Real Charges:           -Rs. {net_fee_oct19:.2f}")
print(f"NET PROFIT/LOSS IN WALLET:   Rs. {(total_gross_pnl_oct19 - net_fee_oct19):.2f}")
